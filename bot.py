import os
import json
import asyncio
from datetime import datetime
from zoneinfo import ZoneInfo
import io

from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup, ReplyKeyboardMarkup, KeyboardButton
from telegram.ext import (
    Application,
    CommandHandler,
    MessageHandler,
    CallbackQueryHandler,
    ContextTypes,
    filters,
)
import openpyxl

# ================= CONFIG =================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

with open(os.path.join(BASE_DIR, "config.txt"), "r", encoding="utf-8") as f:
    BOT_TOKEN = f.read().strip()

DEFAULT_CHANNEL = "@OMORTECH"
DEFAULT_ADMIN_ID = 7815818738

DEFAULT_TARGET_CHAT_IDS = [
    6945991387,
    7815818738,
]

DB_FILE = os.path.join(BASE_DIR, "bot_data.json")
TZ = ZoneInfo("Asia/Dhaka")


# ================= DATABASE & SETTINGS =================

def load_db():
    if not os.path.exists(DB_FILE):
        return {
            "settings": {
                "channel": DEFAULT_CHANNEL,
                "admin_id": DEFAULT_ADMIN_ID,
                "target_chat_ids": DEFAULT_TARGET_CHAT_IDS,
                "maintenance": False,
                "approval_mode": True,
                "min_ids_limit": 1,
                "duplicate_mode": "daily"
            },
            "banned_users": [],
            "pending_submissions": {},
            "days": {}
        }

    try:
        with open(DB_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)

        data.setdefault("settings", {})
        data["settings"].setdefault("channel", DEFAULT_CHANNEL)
        data["settings"].setdefault("admin_id", DEFAULT_ADMIN_ID)
        data["settings"].setdefault("target_chat_ids", DEFAULT_TARGET_CHAT_IDS)
        data["settings"].setdefault("maintenance", False)
        data["settings"].setdefault("approval_mode", True)
        data["settings"].setdefault("min_ids_limit", 1)
        data["settings"].setdefault("duplicate_mode", "daily")
        
        data.setdefault("banned_users", [])
        data.setdefault("pending_submissions", {})
        data.setdefault("days", {})

        return data
    except (json.JSONDecodeError, OSError):
        print("WARNING: Database read failed.")
        return {
            "settings": {
                "channel": DEFAULT_CHANNEL,
                "admin_id": DEFAULT_ADMIN_ID,
                "target_chat_ids": DEFAULT_TARGET_CHAT_IDS,
                "maintenance": False,
                "approval_mode": True,
                "min_ids_limit": 1,
                "duplicate_mode": "daily"
            },
            "banned_users": [],
            "pending_submissions": {},
            "days": {}
        }


DATA = load_db()


def _save_db_file():
    temp_file = DB_FILE + ".tmp"
    with open(temp_file, "w", encoding="utf-8") as f:
        json.dump(DATA, f, ensure_ascii=False, indent=2)
    os.replace(temp_file, DB_FILE)


def save_db():
    asyncio.create_task(asyncio.to_thread(_save_db_file))


def today():
    return datetime.now(TZ).strftime("%Y-%m-%d")


def is_admin(user_id):
    return user_id == DATA["settings"]["admin_id"]


def get_day(date):
    if date not in DATA["days"]:
        DATA["days"][date] = {
            "ids": [],
            "members": {},
        }
    day = DATA["days"][date]
    day.setdefault("ids", [])
    day.setdefault("members", {})
    return day


def get_member(day, user):
    uid = str(user.id)
    members = day["members"]

    if uid not in members:
        members[uid] = {
            "name": user.full_name,
            "received": 0,
            "new": 0,
            "duplicate": 0,
            "rejected": 0,
            "all_ids": [],
            "new_ids": [],
            "duplicate_ids": [],
            "rejected_ids": []
        }

    member = members[uid]
    member["name"] = user.full_name
    member.setdefault("rejected", 0)
    member.setdefault("rejected_ids", [])
    return member


# ================= CHANNEL JOIN & BAN CHECK =================

async def is_joined(user_id, context):
    if is_admin(user_id):
        return True

    channel = DATA["settings"]["channel"]
    try:
        member = await context.bot.get_chat_member(
            chat_id=channel,
            user_id=user_id,
        )
        return member.status in ("member", "administrator", "creator")
    except Exception as e:
        print("Channel membership check failed:", e)
        return False


async def require_join(update, context):
    channel = DATA["settings"]["channel"]
    keyboard = [[
        InlineKeyboardButton(
            "📢 Join Channel",
            url=f"https://t.me/{channel.replace('@', '')}" if channel.startswith("@") else channel,
        )
    ]]

    await update.effective_message.reply_text(
        "🚫 সার্ভিসটি ব্যবহার করতে হলে প্রথমে আমাদের চ্যানেলে জয়েন করতে হবে।\n\n"
        "👉 চ্যানেলে জয়েন করার পর /start দিন।",
        reply_markup=InlineKeyboardMarkup(keyboard),
    )


# ================= START & STATUS =================

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user

    if user.id in DATA["banned_users"]:
        await update.effective_message.reply_text("🚫 আপনাকে সিস্টেম থেকে ব্লক করা হয়েছে।")
        return

    if DATA["settings"]["maintenance"] and not is_admin(user.id):
        await update.effective_message.reply_text("🛠️ বট বর্তমানে মেনটেইনেন্স মোডে আছে। কিছুক্ষণ পর আবার চেষ্টা করুন।")
        return

    if not await is_joined(user.id, context):
        await require_join(update, context)
        return

    date = today()
    day = get_day(date)
    get_member(day, user)
    save_db()

    msg = (
        "<b>🤖 OMOR TECH SUBMISSION BOT</b>\n\n"
        f"👤 <b>Your Name:</b> {user.full_name}\n"
        f"🆔 <b>Your User ID:</b> <code>{user.id}</code>\n\n"
        "🟢 বট সক্রিয় আছে। আপনার প্রয়োজনীয় সংখ্যার ID list অথবা <code>.xlsx</code> ফাইল পাঠান।\n\n"
        "📊 নিজের কাজের হিসাব দেখতে: /status"
    )

    reply_markup = None
    if is_admin(user.id):
        reply_markup = ReplyKeyboardMarkup(
            [[KeyboardButton("👑 অ্যাডমিন প্যানেল")]],
            resize_keyboard=True
        )

    await update.effective_message.reply_text(msg, parse_mode="HTML", reply_markup=reply_markup)


async def status(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user

    if not await is_joined(user.id, context):
        await require_join(update, context)
        return

    uid = str(user.id)
    dates = sorted(DATA["days"].keys(), reverse=True)
    today_date = today()

    lines = [
        "<b>📊 YOUR SUBMISSION STATUS</b>",
        f"👤 <b>Name:</b> {user.full_name}",
        f"🆔 <b>User ID:</b> <code>{user.id}</code>\n",
    ]

    total_received = total_new = total_dup = total_rej = 0
    found = False

    for date in dates:
        member = DATA["days"][date]["members"].get(uid)
        if not member:
            continue

        found = True
        rej_count = member.get("rejected", 0)
        total_received += member["received"]
        total_new += member["new"]
        total_dup += member["duplicate"]
        total_rej += rej_count

        tag = " 🌟 (Today)" if date == today_date else ""
        lines.extend([
            f"📅 <b>Date:</b> <code>{date}</code>{tag}",
            f"  📥 Total: <code>{member['received']}</code> | 🟢 New: <code>{member['new']}</code>",
            f"  ♻️ Dup: <code>{member['duplicate']}</code> | 🚫 Rej: <code>{rej_count}</code>\n",
        ])

    if not found:
        lines.append("📭 আপনি এখনো কোনো আইডি জমা দেননি।")
        await update.effective_message.reply_text("\n".join(lines), parse_mode="HTML")
        return

    lines.extend([
        "━━━━━━━━━━━━━━━━━━",
        "📦 <b>ALL-TIME SUMMARY</b>",
        f"📥 Total Submitted: <code>{total_received}</code>",
        f"🟢 Total New: <code>{total_new}</code>",
        f"♻️ Total Duplicate: <code>{total_dup}</code>",
        f"🚫 Total Rejected: <code>{total_rej}</code>",
    ])

    await update.effective_message.reply_text("\n".join(lines), parse_mode="HTML")


# ================= SEND TO RECEIVERS =================

async def send_receivers(context, text):
    async def _send_to_one(chat_id):
        try:
            for start_pos in range(0, len(text), 3500):
                await context.bot.send_message(
                    chat_id=chat_id,
                    text=text[start_pos:start_pos + 3500],
                    parse_mode="HTML"
                )
        except Exception as e:
            print(f"Message send failed for {chat_id}: {e}")

    await asyncio.gather(*[_send_to_one(cid) for cid in DATA["settings"]["target_chat_ids"]])


# ================= ID & FILE PROCESSING =================

async def process_valid_submission(user, ids, context, update_msg, file_info=None):
    date = today()
    day = get_day(date)
    member = get_member(day, user)

    mode = DATA["settings"]["duplicate_mode"]
    if mode == "all_time":
        all_existing_ids = set()
        for d_data in DATA["days"].values():
            all_existing_ids.update(d_data["ids"])
        seen_today = all_existing_ids
    else:
        seen_today = set(day["ids"])

    seen_message = set()
    new_ids = []
    duplicate_ids = []

    for item in ids:
        if item in seen_today or item in seen_message:
            duplicate_ids.append(item)
        else:
            new_ids.append(item)
        seen_message.add(item)

    day["ids"].extend(ids)
    member["received"] += len(ids)
    member["new"] += len(new_ids)
    member["duplicate"] += len(duplicate_ids)
    member["all_ids"].extend(ids)
    member["new_ids"].extend(new_ids)
    member["duplicate_ids"].extend(duplicate_ids)
    save_db()

    if new_ids:
        single_copy_list = "\n".join([f"<code>{i}</code>" for i in new_ids])
        all_in_one_code = f"<pre>{'\n'.join(new_ids)}</pre>"
        
        report = (
            "<b>📋 NEW APPROVED SUBMISSION</b>\n\n"
            f"👤 <b>Name:</b> {user.full_name}\n"
            f"🆔 <b>User ID:</b> <code>{user.id}</code>\n"
            f"📅 <b>Date:</b> {date}\n\n"
            f"🟢 <b>New IDs ({len(new_ids)}):</b>\n"
            f"{single_copy_list}\n\n"
            f"⚡ <b>COPY ALL IN ONE CLICK:</b>\n"
            f"{all_in_one_code}"
        )

        await send_receivers(context, report)

        if file_info and file_info.get("document_id"):
            for chat_id in DATA["settings"]["target_chat_ids"]:
                try:
                    await context.bot.send_document(
                        chat_id=chat_id,
                        document=file_info["document_id"],
                        caption=f"📁 File from {user.full_name} ({date})"
                    )
                except Exception:
                    pass

    try:
        await context.bot.send_message(
            chat_id=user.id,
            text=(
                f"🎉 <b>আপনার পাঠানো <code>{len(ids)}</code> টি আইডি প্রসেস করা হয়েছে!</b>\n\n"
                f"📥 Total Received: <code>{len(ids)}</code>\n"
                f"🟢 New Approved: <code>{len(new_ids)}</code>\n"
                f"♻️ Duplicate Skipped: <code>{len(duplicate_ids)}</code>"
            ),
            parse_mode="HTML"
        )
    except Exception:
        pass


async def handle_incoming_data(user, ids, context, update_msg, file_info=None):
    min_limit = DATA["settings"]["min_ids_limit"]
    if len(ids) < min_limit:
        await update_msg.reply_text(
            f"⚠️ সর্বনিম্ন <code>{min_limit}</code> টি আইডি একসঙ্গে পাঠাতে হবে। আপনি পাঠিয়েছেন <code>{len(ids)}</code> টি।",
            parse_mode="HTML"
        )
        return

    date = today()
    day = get_day(date)
    mode = DATA["settings"]["duplicate_mode"]
    if mode == "all_time":
        all_existing_ids = set()
        for d_data in DATA["days"].values():
            all_existing_ids.update(d_data["ids"])
        seen_today = all_existing_ids
    else:
        seen_today = set(day["ids"])

    seen_msg = set()
    temp_new = 0
    temp_dup = 0
    for item in ids:
        if item in seen_today or item in seen_msg:
            temp_dup += 1
        else:
            temp_new += 1
        seen_msg.add(item)

    if DATA["settings"]["approval_mode"]:
        sub_id = f"sub_{user.id}_{int(datetime.now().timestamp())}"
        DATA["pending_submissions"][sub_id] = {
            "user_id": user.id,
            "full_name": user.full_name,
            "ids": ids,
            "file_info": file_info
        }
        save_db()

        admin_id = DATA["settings"]["admin_id"]
        admin_preview_name = f"<a href='tg://user?id={user.id}'>{user.full_name}</a>"
        
        preview_text = (
            "<b>🔔 NEW SUBMISSION PENDING</b>\n\n"
            f"👤 <b>User:</b> {admin_preview_name}\n"
            f"🆔 <b>ID:</b> <code>{user.id}</code>\n"
            f"📥 <b>Total IDs:</b> <code>{len(ids)}</code>\n"
            f"🟢 <b>New IDs:</b> <code>{temp_new}</code>\n"
            f"♻️ <b>Duplicate IDs:</b> <code>{temp_dup}</code>\n\n"
            "নিচের বাটন থেকে এক-ক্লিকে প্রসেস করুন:"
        )

        keyboard = [
            [InlineKeyboardButton("🟢 Accept New Only (Skip Dupes)", callback_data=f"acc_{sub_id}")],
            [InlineKeyboardButton("🚫 Reject Entire Batch", callback_data=f"rej_{sub_id}")]
        ]
        
        try:
            await context.bot.send_message(
                chat_id=admin_id,
                text=preview_text,
                reply_markup=InlineKeyboardMarkup(keyboard),
                parse_mode="HTML"
            )
        except Exception as e:
            print("Failed to send notification to admin:", e)

        await update_msg.reply_text(
            f"⏳ আপনার পাঠানো <code>{len(ids)}</code> টি আইডি জমা হয়েছে এবং অ্যাডমিনের অনুমোদনের অপেক্ষায় আছে।",
            parse_mode="HTML"
        )
    else:
        await process_valid_submission(user, ids, context, update_msg, file_info)


async def handle_admin_inputs(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    if not is_admin(user.id):
        return False

    waiting = context.user_data.get("waiting_for")
    if not waiting:
        return False

    text = update.effective_message.text.strip()

    if waiting == "broadcast_message":
        context.user_data["waiting_for"] = None
        all_user_ids = set()
        for d_val in DATA["days"].values():
            for uid in d_val["members"].keys():
                all_user_ids.add(int(uid))

        if not all_user_ids:
            await update.effective_message.reply_text("📭 ব্রডকাস্ট পাঠানোর মতো কোনো রেজিস্টার্ড মেম্বার নেই।")
            return True

        success_count = 0
        fail_count = 0
        broadcast_text = f"📢 <b>ADMIN ANNOUNCEMENT</b>\n\n{text}"

        for uid in all_user_ids:
            try:
                await context.bot.send_message(chat_id=uid, text=broadcast_text, parse_mode="HTML")
                success_count += 1
                await asyncio.sleep(0.05)
            except Exception:
                fail_count += 1

        await update.effective_message.reply_text(
            f"✅ <b>ব্রডকাস্ট সম্পন্ন হয়েছে!</b>\n\n"
            f"🟢 সফলভাবে পাঠানো হয়েছে: <code>{success_count}</code> জনের কাছে\n"
            f"🔴 ব্যর্থ হয়েছে: <code>{fail_count}</code> জন",
            parse_mode="HTML"
        )
        return True

    elif waiting == "msg_user_id":
        parts = text.split(maxsplit=1)
        target_uid_str = parts[0]
        
        if not target_uid_str.isdigit():
            await update.effective_message.reply_text("🔴 সঠিক সংখ্যাসূচক User ID দিন:")
            return True

        target_uid = int(target_uid_str)

        # যদি অ্যাডমিন একই মেসেজে ID এবং কথা লিখে পাঠায় (যেমন: 8114489614 Hi)
        if len(parts) > 1:
            context.user_data["waiting_for"] = None
            msg_text = parts[1]
            try:
                await context.bot.send_message(
                    chat_id=target_uid,
                    text=f"💬 <b>ADMIN MESSAGE:</b>\n\n{msg_text}",
                    parse_mode="HTML"
                )
                await update.effective_message.reply_text(f"✅ সফলভাবে ইউজার <code>{target_uid}</code> এর কাছে আপনার মেসেজটি পাঠানো হয়েছে!", parse_mode="HTML")
            except Exception as e:
                await update.effective_message.reply_text(f"🚫 মেসেজ পাঠানো যায়নি। কারণ: {e}")
            return True
        else:
            context.user_data["target_msg_uid"] = target_uid
            context.user_data["waiting_for"] = "msg_user_text"
            await update.effective_message.reply_text(f"✉️ ইউজার আইডি <code>{target_uid}</code> সিলেক্ট হয়েছে। এবার আপনার মেসেজটি লিখে পাঠান:", parse_mode="HTML")
            return True

    elif waiting == "msg_user_text":
        context.user_data["waiting_for"] = None
        target_uid = context.user_data.get("target_msg_uid")
        if not target_uid:
            await update.effective_message.reply_text("🔴 কোনো ইউজার আইডি সিলেক্ট করা নেই। পুনরায় চেষ্টা করুন।")
            return True

        try:
            await context.bot.send_message(
                chat_id=target_uid,
                text=f"💬 <b>ADMIN MESSAGE:</b>\n\n{text}",
                parse_mode="HTML"
            )
            await update.effective_message.reply_text(f"✅ সফলভাবে ইউজার <code>{target_uid}</code> এর কাছে আপনার মেসেজটি পাঠানো হয়েছে!", parse_mode="HTML")
        except Exception as e:
            await update.effective_message.reply_text(f"🚫 মেসেজ পাঠানো যায়নি। কারণ: {e}")
        return True

    elif waiting == "reject_bad_ids":
        context.user_data["waiting_for"] = None
        raw_ids = [line.strip() for line in text.splitlines() if line.strip().isdigit()]
        if not raw_ids:
            await update.effective_message.reply_text("🚫 কোনো বৈধ আইডি পাওয়া যায়নি।")
            return True

        user_rejected_map = {}
        found_count = 0

        for d_date, d_data in DATA["days"].items():
            for bad_id in raw_ids:
                if bad_id in d_data["ids"]:
                    d_data["ids"].remove(bad_id)
                    found_count += 1

                    for uid, mem in d_data["members"].items():
                        mem.setdefault("rejected", 0)
                        mem.setdefault("rejected_ids", [])

                        if bad_id in mem["all_ids"]:
                            mem["all_ids"].remove(bad_id)
                            
                            if bad_id in mem["new_ids"]:
                                mem["new_ids"].remove(bad_id)
                                mem["new"] = max(0, mem["new"] - 1)
                            elif bad_id in mem["duplicate_ids"]:
                                mem["duplicate_ids"].remove(bad_id)
                                mem["duplicate"] = max(0, mem["duplicate"] - 1)

                            mem["rejected"] += 1
                            mem["rejected_ids"].append(bad_id)

                            user_rejected_map.setdefault(uid, []).append(bad_id)

        save_db()

        for target_uid, bad_list in user_rejected_map.items():
            try:
                formatted_bad_list = "\n".join([f"<code>{i}</code>" for i in bad_list])
                await context.bot.send_message(
                    chat_id=int(target_uid),
                    text=(
                        f"🚫 <b>আপনার পাঠানো নিচের <code>{len(bad_list)}</code> টি আইডি রিজেক্ট/নষ্ট হয়েছে:</b>\n\n"
                        + formatted_bad_list +
                        "\n\n⚠️ আপনার রিজেক্ট স্ট্যাটাসে এগুলো যুক্ত করা হয়েছে।"
                    ),
                    parse_mode="HTML"
                )
            except Exception as e:
                print(f"User Notification Error ({target_uid}):", e)

        await update.effective_message.reply_text(
            f"✅ <b>প্রসেসিং সম্পন্ন হয়েছে!</b>\n\n"
            f"📥 মোট রিজেক্ট আইডি: <code>{len(raw_ids)}</code>\n"
            f"🔍 ডাটাবেজে পাওয়া গেছে: <code>{found_count}</code> টি\n"
            f"👥 সংশ্লিষ্ট মেম্বারদের কাছে মেসেজ পাঠানো হয়েছে এবং রিজেক্ট কাউন্টার আপডেট করা হয়েছে।",
            parse_mode="HTML"
        )
        return True

    elif waiting == "view_user_hist":
        context.user_data["waiting_for"] = None
        target_uid = text
        dates = sorted(DATA["days"].keys(), reverse=True)
        lines = [f"👤 <b>USER HISTORY FOR:</b> <code>{target_uid}</code>\n"]
        found = False

        for d_date in dates:
            mem = DATA["days"][d_date]["members"].get(target_uid)
            if mem:
                found = True
                rej_c = mem.get("rejected", 0)
                lines.append(f"📅 <code>{d_date}</code>\n  📥 Total: <code>{mem['received']}</code> | 🟢 New: <code>{mem['new']}</code>\n  ♻️ Dup: <code>{mem['duplicate']}</code> | 🚫 Rej: <code>{rej_c}</code>\n")

        if found:
            await update.effective_message.reply_text("\n".join(lines), parse_mode="HTML")
        else:
            await update.effective_message.reply_text("🔴 এই ইউজার আইডির কোনো রেকর্ড পাওয়া যায়নি।")
        return True

    elif waiting == "del_hist_type":
        context.user_data["waiting_for"] = None
        input_val = text
        if input_val in DATA["days"]:
            del DATA["days"][input_val]
            save_db()
            await update.effective_message.reply_text(f"🗑️ <code>{input_val}</code> তারিখের সমস্ত হিস্ট্রি সফলভাবে ডিলিট করা হয়েছে।", parse_mode="HTML")
        else:
            deleted = False
            for d_date in DATA["days"]:
                if input_val in DATA["days"][d_date]["members"]:
                    del DATA["days"][d_date]["members"][input_val]
                    deleted = True
            
            if deleted:
                save_db()
                await update.effective_message.reply_text(f"🗑️ ইউজার <code>{input_val}</code> এর সমস্ত হিস্ট্রি ডিলিট করা হয়েছে।", parse_mode="HTML")
            else:
                await update.effective_message.reply_text("🔴 প্রদত্ত তারিখ বা ইউজার আইডি ডাটাবেজে পাওয়া যায়নি।")
        return True

    elif waiting == "add_rcv":
        context.user_data["waiting_for"] = None
        try:
            cid = int(text)
            if cid not in DATA["settings"]["target_chat_ids"]:
                DATA["settings"]["target_chat_ids"].append(cid)
                save_db()
                await update.effective_message.reply_text(f"🟢 সফলভাবে রিসিভার চ্যাট আইডি যোগ করা হয়েছে: <code>{cid}</code>", parse_mode="HTML")
            else:
                await update.effective_message.reply_text("⚠️ এই আইডিটি ইতিমধ্যে লিস্টে আছে।")
        except ValueError:
            await update.effective_message.reply_text("🔴 সঠিক সংখ্যাসূচক চ্যাট আইডি দিন।")
        return True

    elif waiting == "rem_rcv":
        context.user_data["waiting_for"] = None
        try:
            cid = int(text)
            if cid in DATA["settings"]["target_chat_ids"]:
                DATA["settings"]["target_chat_ids"].remove(cid)
                save_db()
                await update.effective_message.reply_text(f"🟢 রিসিভার আইডি রিমুভ করা হয়েছে: <code>{cid}</code>", parse_mode="HTML")
            else:
                await update.effective_message.reply_text("🔴 এই আইডিটি লিস্টে পাওয়া যায়নি।")
        except ValueError:
            await update.effective_message.reply_text("🔴 সঠিক আইডি দিন।")
        return True

    elif waiting == "set_chnl":
        context.user_data["waiting_for"] = None
        DATA["settings"]["channel"] = text
        save_db()
        await update.effective_message.reply_text(f"🟢 নতুন চ্যানেল সেট করা হয়েছে: <code>{text}</code>", parse_mode="HTML")
        return True

    elif waiting == "set_limit":
        context.user_data["waiting_for"] = None
        try:
            limit = int(text)
            DATA["settings"]["min_ids_limit"] = limit
            save_db()
            await update.effective_message.reply_text(f"🟢 মিনিমাম আইডি লিমিট সেট করা হয়েছে: <code>{limit}</code> টি", parse_mode="HTML")
        except ValueError:
            await update.effective_message.reply_text("🔴 সঠিক সংখ্যা দিন।")
        return True

    elif waiting == "ban_user":
        context.user_data["waiting_for"] = None
        try:
            bid = int(text)
            if bid in DATA["banned_users"]:
                DATA["banned_users"].remove(bid)
                save_db()
                await update.effective_message.reply_text(f"🟢 ইউজারকে আনব্যান করা হয়েছে: <code>{bid}</code>", parse_mode="HTML")
            else:
                DATA["banned_users"].append(bid)
                save_db()
                await update.effective_message.reply_text(f"🚫 ইউজারকে ব্যান করা হয়েছে: <code>{bid}</code>", parse_mode="HTML")
        except ValueError:
            await update.effective_message.reply_text("🔴 সঠিক ইউজার আইডি দিন।")
        return True

    return False


async def receive_text(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    if user.id in DATA["banned_users"]:
        return
    if DATA["settings"]["maintenance"] and not is_admin(user.id):
        await update.effective_message.reply_text("🛠️ বট বর্তমানে মেনটেইনেন্স মোডে আছে।")
        return
    if not await is_joined(user.id, context):
        await require_join(update, context)
        return

    text = update.effective_message.text or ""

    if is_admin(user.id) and text == "👑 অ্যাডমিন প্যানেল":
        await admin(update, context)
        return

    # অ্যাডমিন যদি কোনো কমান্ড প্রসেস করতে থাকে তা আগে দেখা হবে
    if is_admin(user.id) and context.user_data.get("waiting_for"):
        handled = await handle_admin_inputs(update, context)
        if handled:
            return

    ids = [line.strip() for line in text.splitlines() if line.strip().isdigit()]
    if not ids:
        return

    await handle_incoming_data(user, ids, context, update.effective_message)


def parse_xlsx_in_thread(file_bytes):
    wb = openpyxl.load_workbook(filename=io.BytesIO(file_bytes), data_only=True)
    sheet = wb.active
    extracted_ids = []
    for row in sheet.iter_rows(values_only=True):
        for cell in row:
            if cell is not None:
                val_str = str(cell).strip()
                if val_str.isdigit():
                    extracted_ids.append(val_str)
    return extracted_ids


async def receive_file(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    if user.id in DATA["banned_users"]:
        return
    if DATA["settings"]["maintenance"] and not is_admin(user.id):
        await update.effective_message.reply_text("🛠️ বট বর্তমানে মেনটেইনেন্স মোডে আছে।")
        return
    if not await is_joined(user.id, context):
        await require_join(update, context)
        return

    document = update.effective_message.document
    if not document:
        return

    filename = document.file_name or ""
    if not filename.lower().endswith(".xlsx"):
        await update.effective_message.reply_text("🚫 শুধু <code>.xlsx</code> ফরম্যাটের ফাইল দিন।", parse_mode="HTML")
        return

    try:
        file = await context.bot.get_file(document.file_id)
        file_bytes = await file.download_as_bytearray()
        
        extracted_ids = await asyncio.to_thread(parse_xlsx_in_thread, file_bytes)

        if extracted_ids:
            file_info = {"document_id": document.file_id}
            await handle_incoming_data(user, extracted_ids, context, update.effective_message, file_info)
        else:
            await update.effective_message.reply_text("🚫 এক্সেল ফাইলের ভেতর কোনো সংখ্যাসূচক আইডি পাওয়া যায়নি।")

    except Exception as e:
        print("XLSX error:", e)
        await update.effective_message.reply_text("🚫 এক্সেল ফাইল রিড করতে সমস্যা হয়েছে।")


# ================= CLEAN SUB-MENU ADMIN PANEL =================

async def admin(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user

    if not is_admin(user.id):
        await update.effective_message.reply_text("🚫 আপনার এই কমান্ডটি ব্যবহার করার অনুমতি নেই।")
        return

    keyboard = [
        [InlineKeyboardButton("📊 Reports & Summary", callback_data="cat_reports")],
        [InlineKeyboardButton("🛠️ ID & Data Management", callback_data="cat_id_mgt")],
        [InlineKeyboardButton("⚙️ Bot System Settings", callback_data="cat_sys_settings")],
        [InlineKeyboardButton("👥 User & Receiver Control", callback_data="cat_user_control")]
    ]

    text = "👑 <b>MAIN ADMIN CONTROL PANEL</b>\n\nযে ক্যাটাগরির কাজ করতে চান নিচের বাটন থেকে সিলেক্ট করুন:"
    
    if update.callback_query:
        await update.callback_query.edit_message_text(text, reply_markup=InlineKeyboardMarkup(keyboard), parse_mode="HTML")
    else:
        await update.effective_message.reply_text(text, reply_markup=InlineKeyboardMarkup(keyboard), parse_mode="HTML")


async def admin_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    user = update.effective_user

    if not is_admin(user.id):
        await query.answer("Access Denied!", show_alert=True)
        return

    await query.answer()
    data = query.data

    # --- MAIN MENU NAVIGATIONS ---
    if data == "cat_main":
        await admin(update, context)
        return

    if data == "cat_reports":
        keyboard = [
            [InlineKeyboardButton("👥 Today Member Summary", callback_data="adm_member_summary")],
            [InlineKeyboardButton("📋 Today All Received IDs", callback_data="adm_all_ids_list")],
            [InlineKeyboardButton("👤 View Specific User History", callback_data="adm_user_hist")],
            [InlineKeyboardButton("👥 All Registered Members", callback_data="adm_all_members")],
            [InlineKeyboardButton("🔙 Back to Main Menu", callback_data="cat_main")]
        ]
        await query.edit_message_text("📊 <b>REPORTS & SUMMARY MENU</b>\n\nপ্রয়োজনীয় রিপোর্ট বেছে নিন:", reply_markup=InlineKeyboardMarkup(keyboard), parse_mode="HTML")
        return

    if data == "cat_id_mgt":
        keyboard = [
            [InlineKeyboardButton("🚫 Reject Specific Bad IDs", callback_data="adm_reject_bad_ids")],
            [InlineKeyboardButton("🗑️ Delete Single Date/User Data", callback_data="adm_del_hist")],
            [InlineKeyboardButton("🔥 Reset All History Data", callback_data="adm_reset_all")],
            [InlineKeyboardButton("🔙 Back to Main Menu", callback_data="cat_main")]
        ]
        await query.edit_message_text("🛠️ <b>ID & DATA MANAGEMENT</b>\n\nআইডি ও ডাটা সংক্রান্ত কাজ করুন:", reply_markup=InlineKeyboardMarkup(keyboard), parse_mode="HTML")
        return

    if data == "cat_sys_settings":
        s = DATA["settings"]
        status_icon = "🔴 OFF" if s["maintenance"] else "🟢 ON"
        approval_icon = "🟢 ON" if s["approval_mode"] else "⚡ Direct"
        mode_icon = "🔄 All-Time" if s["duplicate_mode"] == "all_time" else "📅 Daily"

        keyboard = [
            [InlineKeyboardButton(f"🛡️ Approval Mode: {approval_icon}", callback_data="adm_toggle_approval")],
            [InlineKeyboardButton(f"🛠️ Bot Status: {status_icon}", callback_data="adm_toggle_maint")],
            [InlineKeyboardButton(f"⚙️ Dup Check Mode: {mode_icon}", callback_data="adm_toggle_mode")],
            [InlineKeyboardButton("🔢 Min Limit", callback_data="adm_limit"), InlineKeyboardButton("📢 Channel Link", callback_data="adm_chnl")],
            [InlineKeyboardButton("🔙 Back to Main Menu", callback_data="cat_main")]
        ]
        await query.edit_message_text("⚙️ <b>BOT SYSTEM SETTINGS</b>\n\nবটের গ্লোবাল সেটিংস পরিবর্তন করুন:", reply_markup=InlineKeyboardMarkup(keyboard), parse_mode="HTML")
        return

    if data == "cat_user_control":
        keyboard = [
            [InlineKeyboardButton("➕ Add Receiver", callback_data="adm_add_rcv"), InlineKeyboardButton("➖ Remove Receiver", callback_data="adm_rem_rcv")],
            [InlineKeyboardButton("📢 Broadcast to All Members", callback_data="adm_broadcast")],
            [InlineKeyboardButton("✉️ Message Specific User", callback_data="adm_msg_user")],
            [InlineKeyboardButton("🚫 Ban / Unban User", callback_data="adm_ban")],
            [InlineKeyboardButton("📥 Export Excel Report", callback_data="adm_export")],
            [InlineKeyboardButton("📋 View Full System Config", callback_data="adm_view")],
            [InlineKeyboardButton("🔙 Back to Main Menu", callback_data="cat_main")]
        ]
        await query.edit_message_text("👥 <b>USER & RECEIVER CONTROL</b>\n\nরিসিভার চ্যানেল বা ইউজার নিয়ন্ত্রণ করুন:", reply_markup=InlineKeyboardMarkup(keyboard), parse_mode="HTML")
        return

    # --- ACTIONS & SUB-FUNCTIONS ---
    if data == "adm_member_summary":
        date = today()
        day_data = DATA["days"].get(date)
        
        if not day_data or not day_data["members"]:
            await query.message.reply_text(f"📅 <b>Date:</b> <code>{date}</code>\n\n📭 আজকের দিনে এখনো কোনো আইডি জমা পড়েনি।", parse_mode="HTML")
            return

        lines = [
            "<b>📊 TODAY'S MEMBER SUMMARY</b>",
            f"📅 <b>Date:</b> <code>{date}</code>\n",
            "👥 <b>MEMBER-WISE BREAKDOWN:</b>"
        ]

        total_day_received = 0
        total_day_new = 0
        total_day_dup = 0
        total_day_rej = 0

        for uid, mem in day_data["members"].items():
            adm_rep_name = f"<a href='tg://user?id={uid}'>{mem['name']}</a>"
            total_day_received += mem['received']
            total_day_new += mem['new']
            total_day_dup += mem['duplicate']
            rej_c = mem.get("rejected", 0)
            total_day_rej += rej_c

            lines.extend([
                f"👤 <b>Name:</b> {adm_rep_name}",
                f"  📥 Total: <code>{mem['received']}</code> | 🟢 New: <code>{mem['new']}</code>",
                f"  ♻️ Dup: <code>{mem['duplicate']}</code> | 🚫 Rej: <code>{rej_c}</code>\n"
            ])

        lines.extend([
            "━━━━━━━━━━━━━━━━━━",
            f"📦 <b>Today's Total Received:</b> <code>{total_day_received}</code>",
            f"🟢 <b>Today's Total New:</b> <code>{total_day_new}</code>",
            f"♻️ <b>Today's Total Duplicate:</b> <code>{total_day_dup}</code>",
            f"🚫 <b>Today's Total Rejected:</b> <code>{total_day_rej}</code>"
        ])
        
        await query.message.reply_text("\n".join(lines), parse_mode="HTML")
        return

    if data == "adm_all_ids_list":
        date = today()
        day_data = DATA["days"].get(date)
        
        if not day_data or not day_data["ids"]:
            await query.message.reply_text(f"📅 <b>Date:</b> <code>{date}</code>\n\n📭 আজকের দিনে এখনো কোনো আইডি জমা পড়েনি।", parse_mode="HTML")
            return

        lines = [
            "<b>📋 TODAY'S ALL RECEIVED IDs</b>",
            f"📅 <b>Date:</b> <code>{date}</code>",
            f"📦 <b>Total Count:</b> <code>{len(day_data['ids'])}</code>\n",
            "-----------------------------------"
        ]

        formatted_ids = [f"<code>{i}</code>" for i in day_data["ids"]]
        all_in_one_code = f"<pre>{'\n'.join(day_data['ids'])}</pre>"
        
        full_output = "\n".join(lines) + "\n" + "\n".join(formatted_ids) + f"\n\n⚡ <b>COPY ALL IN ONE CODE:</b>\n{all_in_one_code}"
        
        for start_pos in range(0, len(full_output), 3500):
            await query.message.reply_text(full_output[start_pos:start_pos + 3500], parse_mode="HTML")
        return

    if data == "adm_reject_bad_ids":
        context.user_data["waiting_for"] = "reject_bad_ids"
        await query.message.reply_text(
            "🚫 <b>REJECT SPECIFIC IDs</b>\n\n"
            "যে নষ্ট বা রিজেক্ট করা আইডিগুলো মেম্বারদের ফেরত পাঠাতে চান, সেগুলো এক বা একাধিক লাইনে লিখে মেসেজে পাঠান:",
            parse_mode="HTML"
        )
        return

    if data.startswith("acc_"):
        sub_id = data.replace("acc_", "")
        sub_data = DATA["pending_submissions"].get(sub_id)

        if not sub_data:
            try:
                await query.edit_message_text("⚠️ এই সাবমিশনটি ইতিমধ্যে প্রসেস বা বাতিল করা হয়েছে।")
            except Exception:
                pass
            return

        class DummyUser:
            def __init__(self, uid, name):
                self.id = uid
                self.full_name = name

        d_user = DummyUser(sub_data["user_id"], sub_data["full_name"])
        
        try:
            user_obj = await context.bot.get_chat(sub_data["user_id"])
            if user_obj:
                d_user.full_name = user_obj.full_name or sub_data["full_name"]
        except Exception:
            pass

        await process_valid_submission(
            d_user,
            sub_data["ids"],
            context,
            query.message,
            sub_data["file_info"]
        )

        del DATA["pending_submissions"][sub_id]
        save_db()

        try:
            await query.edit_message_text(f"🟢 Partial Approval Successful!\n👤 User: {sub_data['full_name']} ({len(sub_data['ids'])} IDs Batch)")
        except Exception:
            pass
        return

    if data.startswith("rej_"):
        sub_id = data.replace("rej_", "")
        sub_data = DATA["pending_submissions"].get(sub_id)

        if not sub_data:
            try:
                await query.edit_message_text("⚠️ এই সাবমিশনটি ইতিমধ্যে প্রসেস বা বাতিল করা হয়েছে।")
            except Exception:
                pass
            return

        target_user_id = sub_data["user_id"]
        total_ids_count = len(sub_data["ids"])

        try:
            await context.bot.send_message(
                chat_id=target_user_id,
                text=(
                    f"🚫 <b>আপনার পাঠানো <code>{total_ids_count}</code> টি আইডি গ্রহণযোগ্য হয়নি (Rejected)।</b>\n\n"
                    "অনুগ্রহ করে সঠিক আইডি যাচাই করে পুনরায় চেষ্টা করুন।"
                ),
                parse_mode="HTML"
            )
        except Exception as e:
            print("Failed to send rejection msg to user:", e)

        del DATA["pending_submissions"][sub_id]
        save_db()

        try:
            await query.edit_message_text(f"🚫 Entire Batch Rejected!\n👤 User: {sub_data['full_name']} ({total_ids_count} IDs)")
        except Exception:
            pass
        return

    if data == "adm_all_members":
        all_members = {}
        for d_val in DATA["days"].values():
            for uid, mem in d_val["members"].items():
                if uid not in all_members:
                    all_members[uid] = mem

        if not all_members:
            await query.message.reply_text("📭 ডাটাবেজে এখনো কোনো মেম্বারের রেকর্ড নেই।")
            return

        lines = ["👥 <b>ALL REGISTERED MEMBERS LIST</b>\n(নামে চাপ দিলে সরাসরি প্রোফাইলে নিয়ে যাবে)\n"]
        for uid, mem in all_members.items():
            lines.append(f"👤 Name: <a href='tg://user?id={uid}'>{mem['name']}</a>\n  🆔 User ID: <code>{uid}</code>\n")

        full_output = "\n".join(lines)
        for start_pos in range(0, len(full_output), 3500):
            await query.message.reply_text(full_output[start_pos:start_pos + 3500], parse_mode="HTML")
        return

    if data == "adm_reset_all":
        DATA["days"] = {}
        DATA["pending_submissions"] = {}
        save_db()
        await query.message.reply_text("🔥 <b>সফলভাবে সমস্ত উঠা আইডি ও সাবমিশন হিস্ট্রি রিসেট করা হয়েছে!</b>", parse_mode="HTML")
        return

    elif data == "adm_user_hist":
        context.user_data["waiting_for"] = "view_user_hist"
        await query.message.reply_text("👤 যে ইউজারের হিস্ট্রি দেখতে চান তার <b>User ID</b> লিখে পাঠান:", parse_mode="HTML")

    elif data == "adm_del_hist":
        context.user_data["waiting_for"] = "del_hist_type"
        await query.message.reply_text(
            "🗑️ <b>DELETE HISTORY OPTIONS</b>\n\n"
            "• কোনো নির্দিষ্ট তারিখের হিস্ট্রি মুছতে চাইলে তারিখ লিখুন (যেমন: <code>2026-10-13</code>)\n"
            "• নির্দিষ্ট কোনো ইউজারের অল-টাইম ডাটা মুছতে তার <b>User ID</b> লিখুন:",
            parse_mode="HTML"
        )

    elif data == "adm_broadcast":
        context.user_data["waiting_for"] = "broadcast_message"
        await query.message.reply_text("📢 <b>BROADCAST MESSAGE</b>\n\nসকল রেজিস্টার্ড মেম্বারদের কাছে যে মেসেজটি পাঠাতে চান তা লিখে পাঠান:", parse_mode="HTML")

    elif data == "adm_msg_user":
        context.user_data["waiting_for"] = "msg_user_id"
        await query.message.reply_text("✉️ <b>MESSAGE SPECIFIC USER</b>\n\nযে ইউজারের কাছে মেসেজ পাঠাতে চান তার <b>User ID</b> লিখে পাঠান (অথবা আইডি ও মেসেজ একসাথে পাঠাতে পারেন, যেমন: <code>8114489614 Hello</code>):", parse_mode="HTML")

    if data == "adm_view":
        s = DATA["settings"]
        text = (
            "📋 <b>SYSTEM CONFIGURATION</b>\n\n"
            f"📢 Channel: <code>{s['channel']}</code>\n"
            f"👑 Main Admin ID: <code>{s['admin_id']}</code>\n"
            f"🎯 Receivers: <code>{s['target_chat_ids']}</code>\n"
            f"🔢 Min IDs Limit: <code>{s['min_ids_limit']}</code>\n"
            f"🛡️ Admin Approval Mode: <code>{s['approval_mode']}</code>\n"
            f"🛠️ Maintenance Mode: <code>{s['maintenance']}</code>\n"
            f"⚙️ Duplicate Check: <code>{s['duplicate_mode']}</code>\n"
            f"🚫 Banned Users: <code>{len(DATA['banned_users'])}</code>"
        )
        await query.message.reply_text(text, parse_mode="HTML")

    elif data == "adm_toggle_approval":
        DATA["settings"]["approval_mode"] = not DATA["settings"]["approval_mode"]
        save_db()
        status_text = "ON (অ্যাডমিন অ্যাপ্রুভ লাগবে)" if DATA["settings"]["approval_mode"] else "OFF (সরাসরি ফরোয়ার্ড হবে)"
        await query.message.reply_text(f"🟢 Approval Mode পরিবর্তন করা হয়েছে: {status_text}")

    elif data == "adm_toggle_maint":
        DATA["settings"]["maintenance"] = not DATA["settings"]["maintenance"]
        save_db()
        await query.message.reply_text(f"🟢 Bot Status এখন: {'MAINTENANCE' if DATA['settings']['maintenance'] else 'ACTIVE'}")

    elif data == "adm_toggle_mode":
        DATA["settings"]["duplicate_mode"] = "all_time" if DATA["settings"]["duplicate_mode"] == "daily" else "daily"
        save_db()
        await query.message.reply_text(f"🟢 Duplicate Mode: {DATA['settings']['duplicate_mode'].upper()}")

    elif data == "adm_add_rcv":
        context.user_data["waiting_for"] = "add_rcv"
        await query.message.reply_text("📥 নতুন রিসিভার চ্যাট আইডি (Chat ID) লিখে পাঠান:")

    elif data == "adm_rem_rcv":
        context.user_data["waiting_for"] = "rem_rcv"
        await query.message.reply_text(f"📤 বাদ দেওয়ার জন্য আইডি পাঠান। বর্তমান লিস্ট: <code>{DATA['settings']['target_chat_ids']}</code>", parse_mode="HTML")

    elif data == "adm_chnl":
        context.user_data["waiting_for"] = "set_chnl"
        await query.message.reply_text(f"📢 নতুন প্রাইভেট বা পাবলিক চ্যানেল ইউজারনেম/লিংক পাঠান। বর্তমান: <code>{DATA['settings']['channel']}</code>", parse_mode="HTML")

    elif data == "adm_limit":
        context.user_data["waiting_for"] = "set_limit"
        await query.message.reply_text(f"🔢 মিনিমাম কতটি আইডি দিতে পারবে সেই সংখ্যাটি লিখুন (যেমন: 5)। বর্তমান: <code>{DATA['settings']['min_ids_limit']}</code>", parse_mode="HTML")

    elif data == "adm_ban":
        context.user_data["waiting_for"] = "ban_user"
        await query.message.reply_text("🚫 ব্যান বা আনব্যান করার জন্য ইউজারের User ID পাঠান:")

    elif data == "adm_export":
        def generate_excel():
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Submissions Report"
            ws.append(["Date", "User ID", "Name", "Received", "New", "Duplicate", "Rejected"])

            for d_date, d_val in DATA["days"].items():
                for u_id, mem in d_val["members"].items():
                    ws.append([d_date, u_id, mem["name"], mem["received"], mem["new"], mem["duplicate"], mem.get("rejected", 0)])

            bio = io.BytesIO()
            wb.save(bio)
            bio.seek(0)
            return bio

        excel_bio = await asyncio.to_thread(generate_excel)
        await context.bot.send_document(
            chat_id=user.id,
            document=excel_bio,
            filename=f"report_{today()}.xlsx",
            caption="📊 সম্পূর্ণ সাবমিশনের এক্সেল রিপোর্ট।"
        )


# ================= MAIN =================

async def error_handler(update: object, context: ContextTypes.DEFAULT_TYPE) -> None:
    print(f"⚠️ Network/System Warning: {context.error}")

def main():
    app = (
        Application.builder()
        .token(BOT_TOKEN)
        .read_timeout(30)
        .write_timeout(30)
        .connect_timeout(30)
        .pool_timeout(30)
        .build()
    )

    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("status", status))
    app.add_handler(CommandHandler("admin", admin))
    app.add_handler(CallbackQueryHandler(admin_callback))

    app.add_handler(MessageHandler(filters.Document.ALL, receive_file))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, receive_text))

    app.add_error_handler(error_handler)

    print("🤖 Ultra-Fast High-Tech Telegram Bot is running smoothly...")
    app.run_polling(drop_pending_updates=True)


from keep_alive import keep_alive  # ফাইলের একদম উপরে ইম্পোর্ট করবেন

if __name__ == "__main__":
    keep_alive()  # বট রান হওয়ার আগেই ব্যাকগ্রাউন্ড সার্ভার চালু করবে
    main()